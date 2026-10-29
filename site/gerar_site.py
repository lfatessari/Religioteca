#!/usr/bin/env python3
"""Gera o site Relogioteca (Atlas Horológico) a partir da planilha.

Uso: python3 gerar_site.py <planilha.xlsx> <pasta_saida> [--fotos <pasta_fotos>] [--completo]

  --fotos     pasta do baixar_fotos.py (manifest.json + img/); as fotos baixadas
              são copiadas para <saida>/fotos/ e usadas no lugar dos links
  --completo  gera a página com <!doctype>/<head> (para o GitHub Pages);
              sem isso, sai no formato do Artifact do Claude

Saída:
  <saida>/index.html          página com o índice das marcas embutido
  <saida>/dados/NN.json       fichas completas (história + relógios), em lotes
"""
import json, os, re, sys, unicodedata, urllib.parse
from collections import defaultdict
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from comum import chave, limpar_url, limpo  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
LIMITE_LOTE = 900_000  # bytes por arquivo de dados

# Nomes da aba Relógios que diferem da aba Marcas
APELIDOS = {
    "victorinox": "Victorinox Swiss Army",
    "wylergeneve": "Wyler",
    "zrc": "ZRC 1904",
}


def slug(s):
    s = unicodedata.normalize("NFKD", str(s).lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-") or "marca"


def limpar_site(s):
    s = limpo(s)
    if not s or " " in s.strip():
        return ""
    s = re.sub(r"^https?://", "", s).strip("/")
    return s if "." in s else ""


def limpar_insta(s):
    s = limpo(s)
    if not s:
        return ""
    m = re.search(r"instagram\.com/([A-Za-z0-9._]+)", s)
    if m:
        return "@" + m.group(1)
    s = s.split()[0]
    return s if s.startswith("@") else ("@" + s if re.fullmatch(r"[A-Za-z0-9._]+", s) else "")


def separar_historia(h):
    h = limpo(h)
    aviso = []
    def tira(m):
        aviso.append(m.group(1).strip())
        return ""
    h = re.sub(r"\[\s*⚠️?\s*([^\]]*)\]", tira, h)
    paras = [p.strip() for p in re.split(r"\n\s*\n|\n", h) if p.strip()]
    return paras, "; ".join(a for a in aviso if a)


def lista(s):
    s = limpo(s)
    if not s:
        return []
    itens = [x.strip(" •-\t") for x in re.split(r"\n|;", s)]
    return [x for x in itens if x and len(x) < 120]


def main(planilha, saida, pasta_fotos=None, completo=False, pacotes=False):
    manifest = {}
    if pasta_fotos:
        with open(os.path.join(pasta_fotos, "manifest.json"), encoding="utf-8") as fh:
            manifest = json.load(fh)
    usadas = set()
    wb = openpyxl.load_workbook(planilha, read_only=True, data_only=True)
    faixas = []
    for r in wb["Tabela de faixa de preços"].iter_rows(values_only=True, max_col=3):
        if isinstance(r[0], (int, float)):
            faixas.append((int(r[0]), int(r[1]), int(r[2])))

    def faixa_de(preco):
        for f, a, b in faixas:
            if a <= preco <= b:
                return f
        return 10 if preco > 600000 else 0

    marcas = []
    for r in wb["Marcas"].iter_rows(min_row=2, values_only=True, max_col=13):
        if not limpo(r[0]):
            continue
        nome = limpo(r[0])
        preco = r[5] if isinstance(r[5], (int, float)) else None
        faixa = r[4] if isinstance(r[4], (int, float)) else None
        if faixa is None and isinstance(r[4], str) and r[4].strip().isdigit():
            faixa = int(r[4])
        if faixa is None and preco:
            faixa = faixa_de(preco)
        paras, aviso = separar_historia(r[8])
        marcas.append({
            "marca": nome, "pais": limpo(r[1]), "cidade": limpo(r[2]),
            "carac": limpo(r[3]), "faixa": int(faixa) if faixa else 0,
            "preco": int(preco) if preco else 0,
            "site": limpar_site(r[6]), "insta": limpar_insta(r[7]),
            "hist": paras, "aviso": aviso, "principais": lista(r[9]),
            "relogios": [],
        })

    por_chave = {chave(m["marca"]): m for m in marcas}
    orfas = defaultdict(int)
    vistos = defaultdict(set)
    for r in wb["Relógios"].iter_rows(min_row=2, values_only=True, max_col=5):
        if not limpo(r[0]) or not limpo(r[1]):
            continue
        k = chave(r[0])
        m = por_chave.get(k) or por_chave.get(chave(APELIDOS.get(k, "")))
        if not m:
            orfas[limpo(r[0])] += 1
            continue
        modelo = limpo(r[1])
        fonte = limpar_url(r[2]) or limpar_url(r[3])
        par = (modelo.lower(), fonte)
        if par in vistos[m["marca"]]:
            continue
        vistos[m["marca"]].add(par)
        foto = (manifest.get(fonte) or {}).get("f", "")
        if foto:
            usadas.add(foto)
            foto = "fotos/" + foto
        m["relogios"].append([modelo, foto, limpo(r[4]), fonte])

    # slugs únicos
    usados = set()
    for m in marcas:
        s = slug(m["marca"])
        base, i = s, 2
        while s in usados:
            s, i = f"{base}-{i}", i + 1
        usados.add(s)
        m["id"] = s

    if pacotes and pasta_fotos:
        empacotar(marcas, pasta_fotos, saida)
        usadas.clear()

    # lotes de fichas
    os.makedirs(os.path.join(saida, "dados"), exist_ok=True)
    for f in os.listdir(os.path.join(saida, "dados")):
        os.remove(os.path.join(saida, "dados", f))
    lote, tam, n = {}, 0, 0

    def grava():
        nonlocal lote, tam, n
        if lote:
            with open(os.path.join(saida, "dados", f"{n:02d}.json"), "w", encoding="utf-8") as fh:
                json.dump(lote, fh, ensure_ascii=False, separators=(",", ":"))
            n += 1
        lote, tam = {}, 0

    indice = []
    for m in marcas:
        ficha = {k: m[k] for k in ("hist", "aviso", "principais", "site", "insta", "preco", "relogios")}
        t = len(json.dumps(ficha, ensure_ascii=False).encode())
        if tam and tam + t > LIMITE_LOTE:
            grava()
        lote[m["id"]] = ficha
        tam += t
        busca = " ".join([m["marca"], m["pais"], m["cidade"]] + m["principais"])
        indice.append([m["id"], m["marca"], m["pais"], m["cidade"], m["faixa"],
                       len(m["relogios"]), n, 0 if m["aviso"] else 1, m["carac"],
                       chave_busca(busca)])
    grava()

    with open(os.path.join(AQUI, "modelo.html"), encoding="utf-8") as fh:
        html = fh.read()
    dados = {"faixas": faixas, "marcas": indice}
    html = html.replace("/*__DADOS__*/null", json.dumps(dados, ensure_ascii=False, separators=(",", ":")))
    if completo:
        html = ('<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">'
                '<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">'
                '<meta name="description" content="Atlas das marcas de relógios do @relogioteca: história, faixa de preço e galeria de modelos.">'
                '</head><body>' + html + '</body></html>')
    with open(os.path.join(saida, "index.html"), "w", encoding="utf-8") as fh:
        fh.write(html)

    if pasta_fotos:
        import shutil
        destino = os.path.join(saida, "fotos")
        shutil.rmtree(destino, ignore_errors=True)
        for rel in usadas:
            d = os.path.join(destino, rel)
            os.makedirs(os.path.dirname(d), exist_ok=True)
            shutil.copy2(os.path.join(pasta_fotos, "img", rel), d)

    total = sum(len(m["relogios"]) for m in marcas)
    com = sum(1 for m in marcas if m["relogios"])
    com_foto = sum(1 for m in marcas for r in m["relogios"] if r[1])
    print(f"marcas: {len(marcas)} | com relógios: {com} | relógios: {total} | com foto própria: {com_foto} | lotes: {n}")
    if orfas:
        print("relógios sem marca correspondente:", dict(orfas))


PACOTE_MAX = 750_000   # bytes por pacote de fotos
PREVIA_LADO, PREVIA_QUALIDADE = 300, 50


def _converter(args):
    origem, destino = args
    from PIL import Image
    im = Image.open(origem)
    im.thumbnail((PREVIA_LADO, PREVIA_LADO))
    tmp = destino + ".tmp"
    im.save(tmp, "AVIF", quality=PREVIA_QUALIDADE)
    os.replace(tmp, destino)
    return destino


def empacotar(marcas, pasta_fotos, saida):
    """Prévia no Artifact do Claude: fotos menores (AVIF) juntadas em pacotes .bin.

    O Artifact aceita no máximo 511 arquivos por versão, então as fotos vão em
    pacotes de ~750 KB; a página corta cada foto do pacote quando precisa dela.
    """
    from multiprocessing import Pool
    cache = os.path.join(pasta_fotos, "previa")
    tarefas, pendentes = [], set()
    for m in marcas:
        for r in m["relogios"]:
            if r[1]:
                rel = r[1][len("fotos/"):]
                dest = os.path.join(cache, rel[:-5] + ".avif")
                if dest not in pendentes and not os.path.exists(dest):
                    pendentes.add(dest)
                    os.makedirs(os.path.dirname(dest), exist_ok=True)
                    tarefas.append((os.path.join(pasta_fotos, "img", rel), dest))
    if tarefas:
        print(f"convertendo {len(tarefas)} fotos para a prévia…", flush=True)
        with Pool() as pool:
            for i, _ in enumerate(pool.imap_unordered(_converter, tarefas, chunksize=64), 1):
                if i % 5000 == 0:
                    print(f"  {i}/{len(tarefas)}", flush=True)
    pasta_p = os.path.join(saida, "p")
    os.makedirs(pasta_p, exist_ok=True)
    for f in os.listdir(pasta_p):
        os.remove(os.path.join(pasta_p, f))
    buf, n, total = bytearray(), 0, 0
    ja = {}   # mesma foto usada por mais de um modelo: guarda uma vez só

    def fecha():
        nonlocal buf, n
        if buf:
            with open(os.path.join(pasta_p, f"{n:03d}.bin"), "wb") as fh:
                fh.write(buf)
            n += 1
            buf = bytearray()

    for m in marcas:
        for r in m["relogios"]:
            if not r[1]:
                continue
            rel = r[1][len("fotos/"):]
            if rel in ja:
                r[1] = ja[rel]
                continue
            with open(os.path.join(cache, rel[:-5] + ".avif"), "rb") as fh:
                b = fh.read()
            if len(buf) + len(b) > PACOTE_MAX:
                fecha()
            r[1] = ja[rel] = f"#{n},{len(buf)},{len(b)}"
            buf += b
            total += len(b)
    fecha()
    print(f"prévia: {n} pacotes, {total / 1e6:.0f} MB de fotos")


def chave_busca(s):
    s = unicodedata.normalize("NFKD", s.lower()).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s)


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("planilha")
    ap.add_argument("saida")
    ap.add_argument("--fotos")
    ap.add_argument("--completo", action="store_true")
    ap.add_argument("--pacotes", action="store_true", help="prévia para o Artifact do Claude (fotos em pacotes)")
    a = ap.parse_args()
    main(a.planilha, a.saida, a.fotos, a.completo, a.pacotes)
