import sys, openpyxl, re
from urllib.parse import quote
from openpyxl.worksheet.formula import ArrayFormula
ENT,SAI,MEL,MARCA=sys.argv[1:5]
def enc(u): return quote(u, safe=":/?&=%#+~;,@!$'()*[]-._")
def snap(wb):
    return {ws.title:{c.coordinate:(('ARR',c.value.ref,c.value.text) if isinstance(c.value,ArrayFormula) else c.value) for c in ws._cells.values() if c.value is not None} for ws in wb.worksheets}
wb=openpyxl.load_workbook(ENT); before=snap(wb); w=wb['Relógios']
rows={}
for r in range(2,w.max_row+1):
    if w.cell(r,1).value==MARCA: rows[r]=w.cell(r,2).value
permit=set(); falhas=[]
for l in open(MEL,encoding='utf-8'):
    l=l.rstrip('\n')
    if not l.strip(): continue
    p=l.split('|')+['','']
    nome,img,e,novo=p[0],p[1],p[2],p[3]
    alvo=[r for r,b in rows.items() if b==nome or b.startswith(nome+' - ')]
    if len(alvo)!=1: falhas.append((nome,len(alvo))); continue
    r=alvo[0]
    if img:
        u=enc(img); w.cell(r,3).value=u; w.cell(r,3).hyperlink=u; permit.add(f'C{r}')
    if e:
        old=w.cell(r,5).value or ''
        w.cell(r,5).value=(old.rstrip()+' '+e[1:].strip()).strip() if e.startswith('+') else e; permit.add(f'E{r}')
    if novo:
        b=rows[r]; suf=b[len(nome):]; nb=novo+suf; w.cell(r,2).value=nb; rows[r]=nb; permit.add(f'B{r}')
wb.save(SAI)
after=snap(openpyxl.load_workbook(SAI))
ok=not falhas
for t in before:
    mud={k for k in set(before[t])|set(after[t]) if before[t].get(k)!=after[t].get(k)}
    extra=mud-permit if t=='Relógios' else mud
    print(t,'mudanças:',len(mud),'fora do permitido:',sorted(extra)); ok&=not extra
print('falhas:',falhas); print('RESULTADO:','OK' if ok else 'FALHOU')
