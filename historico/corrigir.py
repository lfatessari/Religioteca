import openpyxl,re,copy,sys
from openpyxl.worksheet.formula import ArrayFormula
SRC,DST=sys.argv[1:3]
L=' - edição limitada, colaboração Peanuts'
ren={
'AV-4139-01':'Flyboy Airmaster 40 Meca-Quartz Peanuts Snoopy Flying Ace Doghouse Green - edição limitada de 250 peças, colaboração Peanuts',
'AV-4150-01':'Spitfire 24 Hour Automatic Peanuts Snoopy Flying Ace Charcoal Grey'+L,
'AV-4150-02':'Spitfire 24 Hour Automatic Peanuts Snoopy Flying Ace Burnt Orange'+L,
'AV-4150-03':'Spitfire 24 Hour Automatic Peanuts Snoopy Flying Ace Azure Blue'+L,
'AV-4161-01':'Flyboy Engineer Automatic Imperial War Museums Fieldmaster Edition Ashen Flight - colaboração Imperial War Museums',
'AV-4161-02':'Flyboy Engineer Automatic Imperial War Museums Fieldmaster Edition Wingman Blue - colaboração Imperial War Museums',
'AV-4161-03':'Flyboy Engineer Automatic Imperial War Museums Fieldmaster Edition Service Steel - colaboração Imperial War Museums',
'AV-4161-04':'Flyboy Engineer Automatic Imperial War Museums Fieldmaster Edition Field Khaki - colaboração Imperial War Museums',
'AV-4088-22':'Hawker Hurricane Carey Dual Time Gutersloh',
'AV-4052-02':'Hawker Hunter Retrograde Chronograph Charcoal Ember Brown',
'AV-4052-05':'Hawker Hunter Retrograde Chronograph Blue Diamond',
'AV-4065-03':'Hawker Harrier II Matador Chronograph Bronze Carbon',
'AV-4149-01':'Spitfire N3200 Imperial War Museums 24 Hour Automatic Eastleigh - edição limitada, colaboração Imperial War Museums',
'AV-4080-02':'Hawker Hunter Duke Chronograph Colerne',
'AV-4011-0U':'Hawker Hurricane Classic Chronograph Grey Black',
'AV-4011-0S':'Hawker Hurricane Classic Chronograph Blood Red',
'AV-4158-01':'Hawker Hurricane Classic Chronograph 1924US Mr Skelly Oak Ash - edição limitada, colaboração 1924US',
'AV-4147-33':'Spitfire Type 300 Automatic Pilot Panel',
'AV-4147-22':'Spitfire Type 300 Automatic Tarmac Tracker',
'AV-4090-06':'Spitfire Smith Automatic Distressed Black',
'AV-4154-01':'Spitfire Smith 24 Hour Automatic Paddy Green - edição limitada (Dia de São Patrício)',
'AV-4152-01':'Spitfire Smith Automatic Snoopy Pilot Error Midnight Maneuver'+L,
'AV-4152-03':'Spitfire Smith Automatic Snoopy Pilot Error Cloudy Confusion'+L,
'AV-4152-04':'Spitfire Smith Automatic Snoopy Pilot Error Deserted Decision'+L,
'AV-4127-01':'Hawker Harrier Pegasus Automatic Piston Black - edição limitada de 500 peças',
'AV-4056-09':'Hawker Harrier II Dual Retrograde Chronograph Admiral Black Steel',
'AV-4077-11':'P-51 Mustang Blakeslee Chronograph Legion',
'AV-4146-01':'P-51 Mustang Hitchcock Automatic 85th Anniversary Steel Ascent - edição limitada',
'AV-4073-11':'Spitfire Type 300 Automatic Oxford Blue',
'AV-4001-03':'Hawker Harrier II Hooper Chronograph Sand',
'AV-4044-02':'Hawker Hurricane Wingman Day Date Blue Bound',
'AV-4044-05':'Hawker Hurricane Wingman Day Date',
'AV-4120-06':'Flyboy Wingman Worldtimer GMT Automatic Islander Viper Black - edição limitada, colaboração Islander (Long Island Watch)',
'AV-4120-07':'Flyboy Wingman Worldtimer GMT Automatic Islander Stealth Blue - edição limitada, colaboração Islander (Long Island Watch)',
'AV-4022-03':'Lancaster Bomber Chronograph',
'AV-4061-01':'Flyboy Centenary 1980s',
}
fam=[(r'^AV-4010-','Hawker Hunter Multifunctional'),(r'^AV-4012-','Hawker Hurricane Traditional'),(r'^AV-4015-','Hawker Hurricane Aviation Chronograph'),(r'^AV-4013-','Hawker Hurricane Chronograph')]
codechg={  # nome antigo (sem código) -> (novo nome, novo código, nota E extra)
'Spitfire Type 300 Automatic Royal British Legion - edição limitada, colaboração Royal British Legion':('Spitfire Type 300 Automatic Royal British Legion Azure Abyss - edição limitada, colaboração Royal British Legion','AV-4073-RBL',None),
'Flyboy Centenary 1980s - mostrador verde, pulseira creme - edição limitada':(None,'AV-4061-03',None),
'Flyboy Lafayette Chronograph - mostrador branco':(None,'AV-4076-11',None),
'Flyboy Lafayette Chronograph - mostrador salmão':(None,'AV-4076-44',None),
'Flyboy Lafayette Chronograph - mostrador azul':(None,'AV-4076-22',' Código deduzido: único código restante da série (-11 branco, -33 preto, -44 salmão).'),
}
rename_code={'AV-4162':('AV-4162-01','Hawker Harrier Pegasus Automatic 60th Anniversary Dunsfold - edição limitada, 60 anos do Hawker Harrier')}
delete_nocode={'Hawker Hunter Retrograde Chronograph Charcoal Ember Brown'}
mel=dict(l.rstrip('\n').split('|',1) for l in open('melhorias_avi8.txt') if l.strip())
def k(c): return re.sub(r'[^A-Z0-9]','',c.upper())
wb=openpyxl.load_workbook(SRC); w=wb['Relógios']
rows=[r for r in range(2,w.max_row+1) if w.cell(r,1).value=='AVI-8']
changed=[]; todel=[]
for r in rows:
    b=w.cell(r,2).value
    if ' - AV-' in b or b.endswith(')') is False and re.search(r' - AV-[\w-]+$',b):
        name,code=b.rsplit(' - ',1)
    else: name,code=b,''
    if not re.match(r'^AV-',code): name,code=b,''
    nb=None
    if code in ren: nb=ren[code]+' - '+code
    for pat,n in fam:
        if re.match(pat,code) and name in ('Hawker Hunter','Hawker Hurricane','Hawker Hurricane - pulseira marrom'): nb=n+' - '+code
    if code in rename_code: nc,nn=rename_code[code]; nb=nn+' - '+nc
    base=name.replace(' (sem referência publicada)','')
    if not code and base in codechg:
        nn,nc,extra=codechg[base]; nb=(nn or base)+' - '+nc
        if extra: w.cell(r,5).value=(w.cell(r,5).value or '')+extra
    if not code and base in delete_nocode: todel.append(r); continue
    if code=='AV-4100-03': w.cell(r,5).value='Cronógrafo, dois fusos. Cor confirmada na coleção oficial.'
    if code=='AV-4090-06': w.cell(r,5).value='Caixa de 42 mm, automático japonês.'
    if nb and nb!=b: w.cell(r,2).value=nb; changed.append((r,b,nb))
    if code and k(code) in mel:
        u=mel[k(code)]; w.cell(r,3).value=u; w.cell(r,3).hyperlink=u; changed.append((r,'C',u[:60]))
# apagar linha duplicada: deslocar para cima as linhas seguintes do bloco (bloco está no fim)
for r in sorted(todel,reverse=True):
    w.delete_rows(r,1)
wb.save(DST)
print(len(changed),'mudanças; apagadas',todel)
for c in changed: print(c)
